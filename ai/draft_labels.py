"""Pre-fill the labelling sheet with CICTT codes suggested by the chat model.

Reads data/labelling_sheet.xlsx, asks gpt-4.1-mini (temperature 0) for one
CICTT code per row using the same definitions a human labeller sees, and
writes data/labelling_sheet_drafted.xlsx with:
  - cictt_code  pre-filled with the suggestion (the reviewer changes it where they disagree)
  - draft_code  the untouched suggestion, kept so agreement can be measured later

The drafted file is not training data. A person reviews it and saves their
corrected copy as data/labelled_seed.xlsx; ai/classify.py reports how often
the draft agreed with that review.

    python ai/draft_labels.py
"""
import json
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from aicommon import PRICE_CHAT_IN, PRICE_CHAT_OUT, ROOT, chat_deployment, openai_client
from make_label_sheet import CODES

SRC = ROOT / "data" / "labelling_sheet.xlsx"
OUT = ROOT / "data" / "labelling_sheet_drafted.xlsx"
WORKERS = 3
CACHE = ROOT / "data" / "staged" / "draft_cache.json"   # resume point, gitignored
VALID = {c for c, _, _ in CODES}
SYSTEM = (
    "You classify aviation accident reports into ICAO CICTT occurrence categories.\n"
    "Pick the ONE category that best describes what happened to the aircraft (the outcome), "
    "not why it happened. If you cannot tell, answer UNK.\n"
    "Output only the code, nothing else.\n\nCategories:\n"
    + "\n".join(f"{c}: {n}. {u}" for c, n, u in CODES)
)


def suggest(client, model, cause: str, findings: str):
    from openai import RateLimitError
    for attempt in range(12):
        try:
            return _suggest(client, model, cause, findings)
        except RateLimitError:
            time.sleep(min(10 * (attempt + 1), 60))      # the deployment's per-minute quota is full
    return _suggest(client, model, cause, findings)


def _suggest(client, model, cause: str, findings: str):
    r = client.chat.completions.create(
        model=model, temperature=0, max_tokens=8,
        messages=[{"role": "system", "content": SYSTEM},
                  {"role": "user", "content": f"Probable cause: {cause}\n\nFindings:\n{findings or '(none)'}"}])
    raw = (r.choices[0].message.content or "").upper()
    found = [t for t in re.findall(r"[A-Z]+(?:-[A-Z]+)?", raw) if t in VALID]
    return (found[0] if found else "UNK"), bool(found), r.usage


def main():
    wb = load_workbook(SRC)
    ws = wb["Label"]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    assert "cictt_code" in col, "unexpected sheet layout"
    draft_idx = ws.max_column + 1
    ws.cell(row=1, column=draft_idx, value="draft_code").font = Font(bold=True)
    ws.cell(row=1, column=draft_idx).fill = PatternFill("solid", fgColor="E2EFDA")
    ws.column_dimensions[ws.cell(row=1, column=draft_idx).column_letter].width = 12

    client, model = openai_client(), chat_deployment()
    rows = list(ws.iter_rows(min_row=2))
    cache = json.loads(CACHE.read_text()) if CACHE.exists() else {}
    todo = [r for r in rows if r[col["record_id"]].value not in cache]
    print(f"{len(rows) - len(todo)} rows already drafted (resuming), {len(todo)} to go", flush=True)
    lock, cost, done, failed = threading.Lock(), [0.0], [0], []

    def work(r):
        return r[col["record_id"]].value, suggest(client, model, r[col["probable_cause"]].value,
                                                  r[col["findings"]].value)
    pool = ThreadPoolExecutor(WORKERS)
    try:
        for fut in as_completed([pool.submit(work, r) for r in todo]):
            try:
                rid, (code, ok, usage) = fut.result()
            except Exception as e:                  # keep every other result; retry this row next run
                failed.append(f"{type(e).__name__}: {str(e)[:120]}")
                continue
            with lock:
                cache[rid] = [code, ok]
                cost[0] += usage.prompt_tokens * PRICE_CHAT_IN / 1e6 + usage.completion_tokens * PRICE_CHAT_OUT / 1e6
                done[0] += 1
                if done[0] % 25 == 0 or done[0] == len(todo):
                    CACHE.write_text(json.dumps(cache))
                    print(f"  drafted {len(rows) - len(todo) + done[0]}/{len(rows)}  (${cost[0]:.3f} this run)", flush=True)
    finally:
        pool.shutdown(wait=False, cancel_futures=True)
        CACHE.write_text(json.dumps(cache))
    if failed:
        print(f"  {len(failed)} rows failed (first error: {failed[0]})")
        raise SystemExit("progress is saved; run the same command again to finish the remaining rows")

    invalid = 0
    for row in rows:
        code, ok = cache[row[col["record_id"]].value]
        invalid += not ok
        row[col["cictt_code"]].value = code
        c = ws.cell(row=row[0].row, column=draft_idx, value=code)
        c.alignment = Alignment(vertical="top")
    how = wb["How to"]
    how.insert_rows(1, 2)
    how["A1"] = ("REVIEW MODE: cictt_code is pre-filled with a model suggestion (also kept in draft_code). "
                 "Read each row and change cictt_code only where you disagree. Do not edit draft_code.")
    how["A1"].font = Font(bold=True)
    wb.save(OUT)
    drafts = [ws.cell(row=r, column=draft_idx).value for r in range(2, ws.max_row + 1)]
    counts = {c: drafts.count(c) for c in sorted(set(drafts), key=lambda x: -drafts.count(x))}
    print(f"wrote {OUT.relative_to(ROOT)}; unusable model answers mapped to UNK: {invalid}")
    print("suggested codes:", ", ".join(f"{c} {k}" for c, k in counts.items()))


if __name__ == "__main__":
    main()
